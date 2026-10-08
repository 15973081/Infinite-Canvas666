import asyncio
import datetime
import io
import zipfile

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from app.services.document_attachments import office_attachment_text
from app.services import media
from app.services.chat import upstream_message_from_record


@pytest.fixture
def docx(tmp_path):
    path = tmp_path/'script.docx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>沉锋第三集：雨夜进入巷子。</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>镜号</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>台词</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:r><w:t>001</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>别回头</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
    return path


@pytest.fixture
def xlsx(tmp_path):
    path = tmp_path/'分镜.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = '分镜'
    sheet.append(['镜号', '场景', '台词', '时长'])
    sheet.append(['001', '雨夜巷口', '别回头|保持警惕\n第二行', 3])
    sheet.append([2, None, None, 4])
    sheet['A3'].number_format = '000'
    sheet.merge_cells('B2:B3')
    sheet.append(['合计', None, None, '=SUM(D2:D3)'])
    sheet.append(['缓存公式', None, None, '=1+2'])
    book.create_sheet('人物').append(['沉锋', datetime.date(2026, 10, 8)])
    buffer = io.BytesIO()
    Image.new('RGB', (2, 2), 'white').save(buffer, format='PNG')
    buffer.seek(0)
    sheet.add_image(SheetImage(buffer), 'F2')
    book.save(path)
    # Excel normally writes cached formula results; openpyxl does not calculate them.
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, 'w') as archive:
        for info, value in entries:
            if info.filename == 'xl/worksheets/sheet1.xml':
                value = value.replace(b'<f>1+2</f><v></v>', b'<f>1+2</f><v>3</v>')
            archive.writestr(info, value)
    return path


def test_docx_preserves_chinese_and_table(docx):
    text = office_attachment_text(str(docx), 100000)
    assert '沉锋第三集：雨夜进入巷子。' in text
    assert '| 001 | 别回头 |' in text


def test_docx_embedded_media_is_not_claimed_as_analyzed(docx):
    with zipfile.ZipFile(docx, 'a') as archive:
        archive.writestr('word/media/image1.png', b'image placeholder')
    assert '文档内嵌图片内容未分析' in office_attachment_text(str(docx), 100000)


def test_docx_missing_optional_dependency_has_install_hint(docx, monkeypatch):
    from markitdown import MarkItDown, FileConversionException, FailedConversionAttempt, MissingDependencyException
    def unavailable(*args, **kwargs):
        raise FileConversionException(attempts=[FailedConversionAttempt(None,
            (MissingDependencyException, MissingDependencyException('mammoth missing'), None))])
    monkeypatch.setattr(MarkItDown, 'convert_local', unavailable)
    with pytest.raises(HTTPException) as error:
        office_attachment_text(str(docx), 100000)
    assert error.value.status_code == 503 and '解析依赖不完整' in error.value.detail


def test_xlsx_preserves_storyboard_structure_and_formula_semantics(xlsx):
    text = office_attachment_text(str(xlsx), 100000)
    assert '## 分镜' in text and '## 人物' in text
    assert '| 2 | 001 | 雨夜巷口 |' in text
    assert '| 3 | 002 | 雨夜巷口 |' in text
    assert '别回头\\|保持警惕<br>第二行' in text
    assert '=SUM(D2:D3)（公式未计算）' in text
    assert '| 5 | 缓存公式 |  |  | 3 |' in text
    assert '2026-10-08' in text
    assert '1 张内嵌图片' in text and '图片内容未分析' in text
    assert 'NaN' not in text


def test_office_and_text_clipping_is_explicit(docx, tmp_path):
    assert '附件内容超过 10 字符' in office_attachment_text(str(docx), 10)
    path = tmp_path/'notes.txt'
    path.write_text('分镜'*40, encoding='utf-8')
    assert '附件内容超过 20 字符' in media.read_text_attachment(str(path), 20)


def test_long_attachment_can_exceed_old_limit(tmp_path):
    path = tmp_path/'notes.txt'
    path.write_text('分镜'*7000+'结尾', encoding='utf-8')
    assert media.read_text_attachment(str(path)).endswith('结尾')


@pytest.mark.parametrize('ext', ['.docx', '.xlsx'])
def test_invalid_office_files_produce_readable_errors(tmp_path, ext):
    path = tmp_path/('broken'+ext)
    path.write_bytes(b'not an Office file')
    with pytest.raises(HTTPException) as error:
        media.read_text_attachment(str(path))
    assert error.value.status_code == 422
    assert '无法解析 Office 附件' in error.value.detail


def test_empty_and_excessive_sheet_ranges_are_rejected(tmp_path):
    path = tmp_path/'empty.xlsx'
    book = Workbook()
    book.save(path)
    with pytest.raises(HTTPException, match='未找到可读取的单元格文字'):
        office_attachment_text(str(path), 100000)
    book.active.cell(250001, 1, 'oversized')
    book.save(path)
    with pytest.raises(HTTPException, match='表格范围过大'):
        office_attachment_text(str(path), 100000)


@pytest.mark.parametrize('fixture_name', ['docx', 'xlsx'])
def test_chat_upstream_receives_document_content(request, fixture_name, monkeypatch):
    path = request.getfixturevalue(fixture_name)
    monkeypatch.setattr(media, 'output_file_from_url', lambda url: str(path))
    message = upstream_message_from_record({'role': 'user', 'content': '分析分镜',
        'attachments': [{'url': '/assets/input/test'+path.suffix, 'kind': 'file', 'name': path.name}]})
    text = message['content'][0]['text']
    assert '分析分镜' in text and '别回头' in text and '附件：'+path.name in text


@pytest.mark.parametrize('endpoint', ['chat', 'chat_stream'])
def test_parse_failure_does_not_save_conversation_or_start_upstream(tmp_path, monkeypatch, endpoint):
    from app.routers import chat as routes
    from app.models import ChatRequest
    from app.services import chat as storage, prompt_skills
    path = tmp_path/'broken.xlsx'
    path.write_bytes(b'broken')
    writes = []
    monkeypatch.setattr(media, 'output_file_from_url', lambda url: str(path))
    monkeypatch.setattr(routes, 'safe_user_id', lambda *args: 'member')
    monkeypatch.setattr(routes, 'prepare_chat', lambda payload: ({'id': 'test'}, '', {}, 'test'))
    monkeypatch.setattr(prompt_skills, 'resolve', lambda ids: [])
    monkeypatch.setattr(storage, 'save_conversation', lambda *args: writes.append(args))
    monkeypatch.setattr(routes, 'save_conversation', lambda *args: writes.append(args))
    monkeypatch.setattr(routes, 'is_apimart_provider', lambda provider: False)
    payload = ChatRequest(message='分析分镜', provider='test', model='test', reference_images=[{
        'url': '/assets/input/broken.xlsx', 'kind': 'file', 'name': '分镜V5.xlsx'}])
    with pytest.raises(HTTPException) as error:
        asyncio.run(getattr(routes, endpoint)(payload, object(), ''))
    assert error.value.status_code == 422
    assert '分镜V5.xlsx' in error.value.detail
    assert writes == []
