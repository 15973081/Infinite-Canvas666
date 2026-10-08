"""Local Office attachments rendered as text for chat, without remote fetching."""
import datetime
import os
import re
import zipfile

from fastapi import HTTPException


OFFICE_EXTENSIONS = {'.docx', '.xlsx'}
MAX_OFFICE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_SHEET_CELLS = 250000


class DocumentContentError(ValueError):
    pass


def clip_attachment_text(text, limit):
    if len(text) > limit:
        return text[:limit] + f'\n\n[附件内容超过 {limit:,} 字符，本次仅提供前 {limit:,} 字符，请拆分文件后分析剩余内容。]'
    return text


def _check_archive(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(entry.file_size for entry in entries) > MAX_OFFICE_EXPANDED_BYTES:
            raise DocumentContentError('文档解压后过大，请拆分文件。')
        return sum(entry.filename.startswith('word/media/') and not entry.is_dir() for entry in entries)


def _cell_text(cell, cached):
    value = cell.value
    if cell.data_type == 'f':
        value = cached.value
        if value is None:
            return f'{cell.value}（公式未计算）'
    if value is None:
        return ''
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat(sep=' ') if isinstance(value, datetime.datetime) else value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int, float)):
        # Preserve common storyboard IDs formatted as 001, 002, etc.
        if re.fullmatch(r'0+', cell.number_format or '') and float(value).is_integer():
            return format(int(value), f'0{len(cell.number_format)}d')
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def _markdown_cell(value):
    return value.replace('\\', '\\\\').replace('|', '\\|').replace('\r\n', '\n').replace('\r', '\n').replace('\n', '<br>')


def _xlsx_markdown(path, limit):
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False)
    try:
        cached_book = load_workbook(path, data_only=True)
        try:
            parts = ['表格说明：合并单元格的内容按覆盖范围重复；没有缓存结果的公式显示原公式，不代表已计算。']
            size = len(parts[0])
            scanned = 0
            populated = False
            for sheet in book.worksheets:
                area = sheet.max_row * sheet.max_column
                scanned += area
                if scanned > MAX_SHEET_CELLS:
                    raise DocumentContentError('表格范围过大，请删除多余空白行列或拆分工作表。')
                merged = {}
                for region in sheet.merged_cells.ranges:
                    anchor = sheet.cell(region.min_row, region.min_col)
                    for row in range(region.min_row, region.max_row+1):
                        for col in range(region.min_col, region.max_col+1):
                            merged[(row, col)] = anchor
                heading = '\n\n## '+sheet.title.replace('\n', ' ')+'\n'
                parts.append(heading)
                size += len(heading)
                header = '| 原始行号 | '+' | '.join(f'列 {i}' for i in range(1, sheet.max_column+1))+' |\n'
                header += '| --- | '+' | '.join('---' for _ in range(sheet.max_column))+' |\n'
                parts.append(header)
                size += len(header)
                for cells in sheet.iter_rows():
                    values = []
                    for cell in cells:
                        source = merged.get((cell.row, cell.column), cell)
                        values.append(_markdown_cell(_cell_text(source, cached_book[sheet.title].cell(source.row, source.column))))
                    if not any(values):
                        continue
                    populated = True
                    line = '| '+str(cells[0].row)+' | '+' | '.join(values)+' |\n'
                    parts.append(line)
                    size += len(line)
                    if size > limit:
                        return ''.join(parts)
                image_count = len(sheet._images)
                if image_count:
                    note = f'\n此工作表包含 {image_count} 张内嵌图片，当前仅提取单元格文字，图片内容未分析。\n'
                    parts.append(note)
                    size += len(note)
            if not populated:
                raise DocumentContentError('未找到可读取的单元格文字；图片分镜请单独上传图片进行分析。')
            return ''.join(parts).strip()
        finally:
            cached_book.close()
    finally:
        book.close()


def office_attachment_text(path, limit):
    """Return bounded Markdown, or a readable send-time error for invalid Office files."""
    try:
        embedded_media = _check_archive(path)
        if os.path.splitext(path)[1].lower() == '.xlsx':
            text = _xlsx_markdown(path, limit)
        else:
            from markitdown import MarkItDown, MissingDependencyException, FileConversionException
            try:
                text = MarkItDown(enable_plugins=False).convert_local(path, keep_data_uris=False).text_content.strip()
            except MissingDependencyException as exc:
                raise HTTPException(503, 'DOCX 解析依赖不完整，请安装 markitdown[docx,xlsx] 并重启服务。') from exc
            except FileConversionException as exc:
                if any(attempt.exc_info and isinstance(attempt.exc_info[1], MissingDependencyException)
                       for attempt in (exc.attempts or [])):
                    raise HTTPException(503, 'DOCX 解析依赖不完整，请安装 markitdown[docx,xlsx] 并重启服务。') from exc
                raise
            if embedded_media and text:
                text += '\n\n文档内嵌图片内容未分析，请单独上传图片进行视觉分析。'
        if not text.strip():
            raise DocumentContentError('未找到可读取的正文或单元格文字。')
        return clip_attachment_text(text, limit)
    except DocumentContentError as exc:
        raise HTTPException(422, str(exc)) from exc
    except ImportError as exc:
        raise HTTPException(503, '文档解析依赖未安装，请安装 requirements.txt 中的 MarkItDown 依赖并重启服务。') from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, '无法解析 Office 附件，请确认文件未损坏、未加密，并保存为 DOCX 或 XLSX 后重新上传。') from exc
