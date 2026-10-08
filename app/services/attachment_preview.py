"""Authoritative local Office preview and immutable, validated model selection."""
import hashlib
import os
import re

from fastapi import HTTPException
from app.config import MAX_ATTACHMENT_TEXT_CHARS
from app.services.document_attachments import (
    _check_archive,_xlsx_markdown,office_attachment_text,clip_attachment_text,DocumentContentError,
    MAX_OFFICE_EXPANDED_BYTES,
)


def digest(path):
    value=hashlib.sha256()
    with open(path,'rb') as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b''): value.update(chunk)
    return value.hexdigest()


def chapters(text):
    """Stable segments in source order, including text before the first heading."""
    pieces=re.split(r'(?m)(?=^#{1,6} )',text)
    return [{'id':'section-'+str(i)+'-'+hashlib.sha256(part.encode()).hexdigest()[:16],
             'title':part.splitlines()[0].lstrip('# ').strip()[:160] or '正文','text':part}
            for i,part in enumerate(pieces) if part.strip()]


def extract(path,selection=None,limit=MAX_ATTACHMENT_TEXT_CHARS):
    if not path or not os.path.isfile(path): raise HTTPException(404,'附件不存在。')
    kind=os.path.splitext(path)[1].lower().lstrip('.')
    if kind not in {'docx','xlsx'}: raise HTTPException(415,'解析预览和范围选择仅支持 DOCX、XLSX。')
    sha256=digest(path)
    if selection and (selection['kind']!=kind or selection['sha256']!=sha256):
        raise HTTPException(409,'附件已变化，请重新预览并选择范围。')
    try:
        images=_check_archive(path)
        warnings=['仅提取正文和单元格文字；内嵌图片、图表、批注及修订的完整视觉效果未分析。']
        result={'kind':kind,'sha256':sha256,'sections':[],'sheets':[],'warnings':warnings}
        if kind=='docx':
            text=office_attachment_text(path,MAX_OFFICE_EXPANDED_BYTES)
            sections=chapters(text)
            result['sections']=[{k:s[k] for k in ('id','title')} for s in sections]
            if selection:
                chosen=selection.get('sections',[])
                if (selection.get('sheets') or not chosen or len(set(chosen))!=len(chosen)
                    or set(chosen)-{s['id'] for s in sections}):
                    raise DocumentContentError('请选择有效且不重复的 DOCX 章节。')
                text='\n\n'.join(s['text'] for s in sections if s['id'] in chosen)
                if images: text+='\n\n文档内嵌图片内容未分析，请单独上传图片进行视觉分析。'
        else:
            from openpyxl import load_workbook
            book=load_workbook(path,data_only=False)
            try:
                result['sheets']=[{'name':s.title,'rows':s.max_row,'columns':s.max_column,'hidden':s.sheet_state!='visible',
                                  'images':len(s._images)} for s in book.worksheets]
            finally: book.close()
            warnings.append('公式使用文件中的缓存结果；无缓存时显示原公式，不执行计算。合并单元格按覆盖区域显示锚点文字。')
            if any(s['hidden'] for s in result['sheets']): warnings.append('文件包含隐藏工作表；全文引用也会包含隐藏表，请核对选择范围。')
            if selection and (selection.get('sections') or not selection.get('sheets')):
                raise DocumentContentError('请选择有效的 XLSX 工作表和区域。')
            try: text=_xlsx_markdown(path,limit,selection.get('sheets') if selection else None)
            except DocumentContentError as exc:
                if selection: raise
                result['warnings'].append(str(exc)); text=''
        result['text']=clip_attachment_text(text,limit)
        result['truncated']=len(text)>limit
        return result
    except DocumentContentError as exc: raise HTTPException(422,str(exc)) from exc
    except HTTPException: raise
    except ImportError as exc: raise HTTPException(503,'文档解析依赖不完整，请安装 requirements.txt 并重启服务。') from exc
    except Exception as exc: raise HTTPException(422,'无法解析 Office 附件，请确认文件未损坏或加密。') from exc


def selected_text(path,selection,limit):
    result=extract(path,selection,limit)
    if not result['text']: raise HTTPException(422,'选择范围没有可读取的文字。')
    return result['text']


def preview(ref):
    from app.services.media import output_file_from_url
    result=extract(output_file_from_url(ref['url']),ref.get('selection'))
    text=result.pop('text')
    return {**result,'text':text[:20000],'characters':len(text),'preview_truncated':len(text)>20000}
